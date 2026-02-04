"""excel_to_sql.py

Convert an Excel sheet to a SQLite3 SQL script.

Usage:
    python excel_to_sql.py input.xlsx output.sql Sheet1 dlt645_dataitem_A1

Defaults: input.xlsx=data.xlsx, output.sql=data.sql, sheet=Sheet1, table=dlt645_dataitem_A1
"""
import sqlite3
import sys
import re
from openpyxl import load_workbook


def sanitize_col(name: str, index: int):
    if not name:
        return f"col_{index}"
    s = str(name).strip()
    # replace spaces with underscore
    s = re.sub(r"\s+", "_", s)
    # keep only word characters (letters, digits, underscore)
    s = re.sub(r"[^\w]", "", s)
    if not s:
        return f"col_{index}"
    # if it starts with digit, prefix with c_
    if re.match(r"^\d", s):
        s = "c_" + s
    return s


def quote_value(conn, v):
    if v is None:
        return "NULL"
    if isinstance(v, (int, float)):
        return str(v)
    # use sqlite quote() for proper escaping
    cur = conn.execute("select quote(?)", (str(v),))
    return cur.fetchone()[0]


def excel_to_sql(input_path: str = 'data.xlsx', output_path: str = 'data.sql', sheet: str = 'Sheet1', table: str = 'dlt645_dataitem_A1'):
    wb = load_workbook(input_path, data_only=True)
    if sheet not in wb.sheetnames:
        raise ValueError(f"Sheet '{sheet}' not found in {input_path}")
    ws = wb[sheet]

    conn = sqlite3.connect(':memory:')

    # read header (first row) as column names
    header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    cols = [sanitize_col(header[i], i + 1) for i in range(len(header))]

    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(f"DROP TABLE IF EXISTS \"{table}\";\n")
        f.write(f"CREATE TABLE \"{table}\" (\n")
        col_defs = [f'  "{c}" TEXT' for c in cols]
        f.write(',\n'.join(col_defs))
        f.write('\n);\n\n')

        col_list = ', '.join([f'"{c}"' for c in cols])

        # iterate data rows (starting from row 2)
        hex_cols = {"DI3", "DI2", "DI1", "DI0"}
        for row in ws.iter_rows(min_row=2, values_only=True):
            vals = []
            for i, v in enumerate(row[:len(cols)]):
                colname = cols[i]
                if colname in hex_cols and v is not None:
                    # normalize value to 2-digit uppercase hex (preserve as TEXT)
                    s = str(v).strip()
                    try:
                        iv = int(s, 16)
                        hexs = format(iv, '02X')
                        cur = conn.execute("select quote(?)", (hexs,))
                        vals.append(cur.fetchone()[0])
                        continue
                    except Exception:
                        # fallback to default quoting
                        pass

                vals.append(quote_value(conn, v))

            f.write(f'INSERT INTO "{table}" ({col_list}) VALUES ({", ".join(vals)});\n')

    conn.close()
    print(f"Wrote SQL to {output_path}")


if __name__ == '__main__':
    inp = sys.argv[1] if len(sys.argv) > 1 else 'data.xlsx'
    out = sys.argv[2] if len(sys.argv) > 2 else 'data.sql'
    sheet = sys.argv[3] if len(sys.argv) > 3 else 'Sheet1'
    table = sys.argv[4] if len(sys.argv) > 4 else 'dlt645_dataitem_A1'
    excel_to_sql(inp, out, sheet, table)
