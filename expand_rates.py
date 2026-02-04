from openpyxl import load_workbook
import sys

ELLIPSIS = '…'


def try_parse_int(value):
    if value is None:
        return None, None, None
    s = str(value).strip()
    # detect negative
    neg = s.startswith('-')
    s_n = s[1:] if neg else s
    # prefer hex parsing (user expects hex sequences like 01,0A,3F)
    try:
        v = int(s_n, 16)
        return v, 16, {'width': len(s_n), 'neg': neg, 'orig': s, 'hex_upper': any(c.isalpha() and c.isupper() for c in s_n)}
    except Exception:
        pass
    # fallback decimal
    try:
        v = int(s_n, 10)
        return v, 10, {'width': len(s_n), 'neg': neg, 'orig': s}
    except Exception:
        return None, None, None


def format_int(v, base, meta):
    if v is None:
        return None
    if base == 10:
        s = str(v)
        width = meta.get('width') if meta else None
        if width:
            s = s.zfill(width)
        if meta and meta.get('neg'):
            s = '-' + s
        return s
    elif base == 16:
        # use uppercase hex and preserve width when available (default to 2)
        s = format(v, 'X')
        width = meta.get('width') if meta else 2
        if width:
            s = s.zfill(width)
        return s
    else:
        return str(v)


def is_all_ellipsis(row_vals):
    # True if every non-empty cell is ellipsis and at least one ellipsis present
    has_ellipsis = False
    any_nonempty = False
    for v in row_vals:
        if v is None:
            continue
        any_nonempty = True
        if str(v).strip() == ELLIPSIS:
            has_ellipsis = True
        else:
            return False
    return any_nonempty and has_ellipsis


def find_prev_non_ellipsis(ws, start_row):
    r = start_row - 1
    while r >= 1:
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        if not is_all_ellipsis(row_vals):
            return r
        r -= 1
    return None


def expand_ellipsis_rows(file_path: str, sheet_name: str, output_path: str = None):
    if output_path is None:
        output_path = file_path.replace('.xlsx', '') + '_expanded.xlsx'

    wb = load_workbook(file_path)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet '{sheet_name}' not found")
    ws = wb[sheet_name]

    # We'll expand ranges for the sequence column (default column 3 -> C) where a concrete
    # value appears, followed by one or more ellipsis rows, then another concrete value.
    # The ellipsis block will be replaced with the full sequence of hex values between
    # the two concrete values (exclusive), copying other columns from the start row.
    seq_col = 3
    processed_blocks = 0
    r = 1
    max_col = ws.max_column
    while r <= ws.max_row:
        start_val = ws.cell(row=r, column=seq_col).value
        if start_val is None or str(start_val).strip() == ELLIPSIS:
            r += 1
            continue

        # find next row with a concrete value in seq_col
        t = r + 1
        while t <= ws.max_row and (ws.cell(row=t, column=seq_col).value is None or str(ws.cell(row=t, column=seq_col).value).strip() == ELLIPSIS):
            t += 1

        if t > ws.max_row:
            break

        # only consider blocks where there are intermediate rows to replace
        if t == r + 1:
            r = t
            continue

        # check whether we should skip (if all intermediate rows are full-ellipsis rows)
        any_non_all = False
        for k in range(r + 1, t):
            row_vals = [ws.cell(row=k, column=c).value for c in range(1, max_col + 1)]
            if not is_all_ellipsis(row_vals):
                any_non_all = True
                break
        if not any_non_all:
            # skip expansion for this block
            r = t
            continue

        # parse start and end as integers (hex preferred)
        pv, pbase, pmeta = try_parse_int(ws.cell(row=r, column=seq_col).value)
        nv, nbase, nmeta = try_parse_int(ws.cell(row=t, column=seq_col).value)
        if pv is None or nv is None:
            r = t
            continue

        if nv <= pv + 1:
            # nothing to insert (adjacent or no gap)
            r = t
            continue

        seq = list(range(pv + 1, nv))
        num_existing = t - r - 1
        num_to_insert = len(seq)

        # remove existing intermediate rows
        ws.delete_rows(r + 1, num_existing)
        # insert the required rows
        ws.insert_rows(r + 1, num_to_insert)

        # fill inserted rows: copy other columns from start row; seq values into seq_col
        for i, v in enumerate(seq):
            target = r + 1 + i
            for c in range(1, max_col + 1):
                if c == seq_col:
                    ws.cell(row=target, column=c).value = format_int(v, pbase, pmeta)
                else:
                        # for description column, try to replace the fee-rate number inside text
                        desc_col = 10
                        src_val = ws.cell(row=r, column=c).value
                        if c == desc_col:
                            # choose a template: prefer start row's description, else end row's
                            template = src_val
                            if template is None or (isinstance(template, str) and template.strip() == ELLIPSIS):
                                template = ws.cell(row=t, column=c).value

                            new_desc = template
                            if isinstance(template, str) and template:
                                import re

                                def replace_fee_number(text, new_number):
                                    # prefer patterns like '费率 <num>' then fallback to first standalone number
                                    m = re.search(r'(费率)\s*(\d+)', text)
                                    if m:
                                        start, end = m.span(2)
                                        return text[:start] + str(new_number) + text[end:]
                                    # fallback: find first number after '费' word
                                    m2 = re.search(r'费[^\d]*(\d+)', text)
                                    if m2:
                                        start, end = m2.span(1)
                                        return text[:start] + str(new_number) + text[end:]
                                    # fallback: replace first standalone number
                                    m3 = re.search(r'\b(\d+)\b', text)
                                    if m3:
                                        start, end = m3.span(1)
                                        return text[:start] + str(new_number) + text[end:]
                                    return text

                                new_desc = replace_fee_number(template, v)
                            ws.cell(row=target, column=c).value = new_desc
                        else:
                            ws.cell(row=target, column=c).value = ws.cell(row=r, column=c).value

        processed_blocks += 1
        # move r to the row after the inserted block (which now sits before the original t row)
        r = r + num_to_insert + 1

    wb.save(output_path)
    print(f"Processed {processed_blocks} blocks; saved to {output_path}")


if __name__ == '__main__':
    file_path = 'data.xlsx'
    sheet_name = 'Sheet1'
    out = 'data_expanded.xlsx'
    expand_ellipsis_rows(file_path, sheet_name, out)
